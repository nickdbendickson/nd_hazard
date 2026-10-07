"""Tests of the Word submittal report (core/report_docx.py, algorithms_report.py).

The dam folder used here comes from the real workflow (`workflow.run_dam` on the
synthetic valley) with only the LISFLOOD-FP executable stubbed, so the report
reads exactly what a run writes. Extras written by other modules (sensitivity,
consequences, manifest, cascade) are fabricated beside it."""
import atexit
import base64
import json
import os
import re
import shutil
import struct
import sys
import tempfile
import xml.dom.minidom as minidom
import zipfile
import zlib

import numpy as np

from nd_breach_runner.tests.registry import test
from nd_breach_runner.core import report_docx as R

_JPEG_8x4 = base64.b64decode(
    "/9j/4AAQSkZJRgABAQEAYABgAAD/2wBDABALDA4MChAODQ4SERATGCgaGBYWGDEjJR0oOjM9PDkzODdASFxOQERXRTc4"
    "UG1RV19iZ2hnPk1xeXBkeFxlZ2P/2wBDARESEhgVGC8aGi9jQjhCY2NjY2NjY2NjY2NjY2NjY2NjY2NjY2NjY2NjY2Nj"
    "Y2NjY2NjY2NjY2NjY2NjY2NjY2P/wAARCAAEAAgDASIAAhEBAxEB/8QAHwAAAQUBAQEBAQEAAAAAAAAAAAECAwQFBgcI"
    "CQoL/8QAtRAAAgEDAwIEAwUFBAQAAAF9AQIDAAQRBRIhMUEGE1FhByJxFDKBkaEII0KxwRVS0fAkM2JyggkKFhcYGRol"
    "JicoKSo0NTY3ODk6Q0RFRkdISUpTVFVWV1hZWmNkZWZnaGlqc3R1dnd4eXqDhIWGh4iJipKTlJWWl5iZmqKjpKWmp6ip"
    "qrKztLW2t7i5usLDxMXGx8jJytLT1NXW19jZ2uHi4+Tl5ufo6erx8vP09fb3+Pn6/8QAHwEAAwEBAQEBAQEBAQAAAAAA"
    "AAECAwQFBgcICQoL/8QAtREAAgECBAQDBAcFBAQAAQJ3AAECAxEEBSExBhJBUQdhcRMiMoEIFEKRobHBCSMzUvAVYnLR"
    "ChYkNOEl8RcYGRomJygpKjU2Nzg5OkNERUZHSElKU1RVVldYWVpjZGVmZ2hpanN0dXZ3eHl6goOEhYaHiImKkpOUlZaX"
    "mJmaoqOkpaanqKmqsrO0tba3uLm6wsPExcbHyMnK0tPU1dbX2Nna4uPk5ebn6Onq8vP09fb3+Pn6/9oADAMBAAIRAxEA"
    "PwDHooorhPqD/9k=")


def _png(w, h, dpi=None, rgb=(200, 40, 40)):
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(t, d):
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    out = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
    if dpi:
        ppm = int(round(dpi / 0.0254))
        out += chunk(b"pHYs", struct.pack(">IIB", ppm, ppm, 1))
    return out + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def _tmp():
    d = tempfile.mkdtemp(prefix="ndbr_test_")
    atexit.register(shutil.rmtree, d, True)
    return d


def _check_package(path):
    """Structural checks that hold for every file this module writes; returns the
    parsed document.xml text and the zip's names."""
    z = zipfile.ZipFile(path)
    assert z.testzip() is None
    names = set(z.namelist())
    assert z.namelist()[0] == "[Content_Types].xml"
    for part in ("[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml",
                 "word/numbering.xml", "word/settings.xml", "word/_rels/document.xml.rels",
                 "docProps/core.xml", "docProps/app.xml"):
        assert part in names, part
    for n in names:                        # XML well-formed for every part
        if n.endswith((".xml", ".rels")):
            minidom.parseString(z.read(n))
    ct = z.read("[Content_Types].xml").decode()
    for n in names:
        if n.endswith(".xml") and n != "[Content_Types].xml":
            assert '"/{}"'.format(n) in ct, "no content type for " + n
    rels = z.read("word/_rels/document.xml.rels").decode()
    ids = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    ids.update({a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    for rid, target in ids.items():
        assert "word/" + target in names, (rid, target)
    doc = z.read("word/document.xml").decode()
    for rid in re.findall(r'r:(?:embed|id)="(rId\d+)"', doc):
        assert rid in ids, rid
    marks = re.findall(r'<w:bookmarkStart w:id="(\d+)" w:name="([^"]+)"', doc)
    assert len({m[0] for m in marks}) == len(marks) and len({m[1] for m in marks}) == len(marks)
    for anchor in re.findall(r'w:anchor="([^"]+)"', doc):
        assert anchor in {m[1] for m in marks}, anchor
    pr = re.findall(r'<wp:docPr id="(\d+)"', doc)
    assert len(set(pr)) == len(pr)
    styles = z.read("word/styles.xml").decode()
    sids = set(re.findall(r'w:styleId="([^"]+)"', styles))
    for s in set(re.findall(r'<w:(?:pStyle|rStyle|tblStyle) w:val="([^"]+)"', doc)):
        assert s in sids, "undefined style " + s
    nums = set(re.findall(r'<w:num w:numId="(\d+)"', z.read("word/numbering.xml").decode()))
    for n in set(re.findall(r'<w:numId w:val="(\d+)"', doc)):
        assert n in nums, "undefined numId " + n
    return doc, names


# ----------------------------------------------------------------------------- the builder
@test
def docx_builder_writes_a_complete_valid_package():
    d = _tmp()
    doc = R.Docx(title="A & B <report>", header_left="Left", header_right="Right",
                 footer_left="Foot", number_headings=True)
    doc.title_block("Title", "Sub", [("Engineer", "E. Ngineer"), ("Empty", "")],
                    logo=_png(40, 20), note="note")
    doc.toc()
    doc.page_break()
    doc.heading("First", 1)
    doc.heading("Second", 2)
    doc.heading("Third", 3)
    doc.heading("Another", 1)
    doc.paragraph("plain **bold** *it* `code` snake_case_name", markup=True)
    doc.paragraph([{"text": "a", "b": True}, {"text": "b", "i": True}, "c"])
    doc.bullets(["one", ("two", ["two.a", "two.b"])])
    doc.bullets(["x", "y"], ordered=True)
    doc.bullets(["again"], ordered=True)
    doc.table([["H1", "H2", "H3"], ["a", "1", {"text": "shaded", "shade": "FFF2CC"}],
               ["b", "2", "long text " * 20]], widths=[1, 1, 3], align=[None, "right", None],
              caption="A table")
    doc.image(_png(300, 150, dpi=150), caption="A figure")
    doc.callout("careful", kind="warn")
    doc.code_block("x = 1\ny = 2")
    path = doc.save(os.path.join(d, "a.docx"))
    xml, names = _check_package(path)
    z = zipfile.ZipFile(path)
    core = z.read("docProps/core.xml").decode()
    assert "<dc:title>A &amp; B &lt;report&gt;</dc:title>" in core
    assert "<dc:creator></dc:creator>" in core              # author blank by default
    assert "word/header1.xml" in names and "word/footer1.xml" in names
    assert "word/header2.xml" in names                       # cover page has none
    foot = z.read("word/footer1.xml").decode()
    assert "PAGE" in foot and "NUMPAGES" in foot and "Foot" in foot
    assert 'TOC \\o "1-3"' in xml
    assert "<w:titlePg/>" in xml
    assert 'updateFields' in z.read("word/settings.xml").decode()
    info = R.inspect_docx(path)
    assert [h[0] for h in info["headings"]] == [1, 2, 3, 1], info["headings"]
    assert info["headings"][1][1] == "1.1 Second" and info["headings"][3][1] == "2 Another"
    assert info["tables"] == 2 and info["images"] == 2      # the logo and the figure
    # the TOC field's fallback lists every heading as a plain, linked entry
    for h in info["headings"]:
        assert h[1] in info["text"]
    assert xml.count('w:fldCharType="begin"') == 3          # TOC, and SEQ for the table and figure
    assert "Figure " in info["text"] and "Table " in info["text"]
    # header row: shaded, repeated; widths sum to the text width
    assert "<w:tblHeader/>" in xml and 'w:fill="D9E2F3"' in xml
    widths = [int(w) for w in re.findall(r'<w:gridCol w:w="(\d+)"', xml)][-3:]
    assert sum(widths) == 9360, widths
    # the two numbered lists restart
    assert len(set(re.findall(r'<w:numId w:val="(\d+)"/>', xml)) - {"1"}) == 2
    # a document without a header, footer or page numbers still packages
    p2 = R.Docx(page_numbers=False)
    p2.paragraph("x")
    p2.save(os.path.join(d, "b.docx"))
    _doc2, n2 = _check_package(os.path.join(d, "b.docx"))
    assert "word/footer1.xml" not in n2
    # A4 and landscape change the page
    p3 = R.Docx(page="a4", landscape=True)
    p3.paragraph("x")
    p3.save(os.path.join(d, "c.docx"))
    x3, _ = _check_package(os.path.join(d, "c.docx"))
    assert 'w:w="16838" w:h="11906" w:orient="landscape"' in x3
    shutil.rmtree(d)


@test
def docx_images_get_correct_emu_sizes_and_are_deduplicated():
    assert R.image_info(_png(300, 150, dpi=150))[:3] == ("png", 300, 150)
    fmt, w, h, dx, dy = R.image_info(_JPEG_8x4)
    assert (fmt, w, h) == ("jpeg", 8, 4) and abs(dx - 96) < 1 and abs(dy - 96) < 1
    d = _tmp()
    doc = R.Docx()
    a = doc.image(_png(300, 150, dpi=150))                 # native: 2 in x 1 in
    assert abs(a[0] - 2.0) < 1e-9 and abs(a[1] - 1.0) < 1e-9
    doc.image(_png(300, 150, dpi=150), width_in=3.0)       # aspect kept
    doc.image(_JPEG_8x4, width_in=3.0)
    doc.image(_png(1000, 1000, dpi=72))                    # 13.9 in: shrunk to the text width
    doc.image(_png(10, 10))                                # no dpi: 96
    path = doc.save(os.path.join(d, "i.docx"))
    xml, names = _check_package(path)
    ext = [(int(a), int(b)) for a, b in re.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"/>', xml)]
    assert ext[0] == (1828800, 914400)
    assert ext[1] == (2743200, 1371600)
    assert ext[2] == (2743200, 1371600)                    # 8 x 4 px at 3 in wide
    assert ext[3][0] == ext[3][1] and ext[3][0] <= int(6.5 * 914400)
    assert abs(ext[4][0] - 914400 * 10 / 96.0) <= 1
    assert xml.count("<a:ext ") == 5
    media = [n for n in names if n.startswith("word/media/")]
    assert len(media) == 4, media                           # the two identical PNGs share one
    assert any(n.endswith(".jpeg") for n in media)
    for bad in (b"GIF89a....", b"", b"\x89PNG\r\n\x1a\nxx", b"\xff\xd8\xff\xd9"):
        try:
            R.image_info(bad)
            assert False, bad
        except ValueError:
            pass
    try:
        R.Docx().image(b"not an image")
        assert False
    except ValueError:
        pass
    shutil.rmtree(d)


@test
def docx_text_is_escaped_and_control_characters_removed():
    d = _tmp()
    doc = R.Docx(title="t\x00itle")
    nasty = u"a<b>&\"'\x00\x01\x0b\x1f\ud800 end￾ é中 \tTAB\nLINE"
    doc.heading(nasty, 1)
    doc.paragraph(nasty)
    doc.table([[nasty, None, 5, True]])
    doc.bullets([nasty])
    path = doc.save(os.path.join(d, "n.docx"))
    xml, _ = _check_package(path)
    assert "\x00" not in xml and u"\ud800" not in xml
    assert u"a&lt;b&gt;&amp;" in xml and u"é中" in xml
    assert "<w:tab/>" in xml and "<w:br/>" in xml
    shutil.rmtree(d)


@test
def docx_loads_in_python_docx_when_it_is_installed():
    try:
        import docx as pydocx
    except ImportError:
        return
    d = _tmp()
    doc = R.Docx(title="T", number_headings=False)
    doc.heading("H1", 1)
    doc.heading("H2", 2)
    doc.paragraph("body")
    doc.table([["a", "b"], ["1", "2"]])
    doc.table([["c"], ["3"]])
    doc.image(_png(30, 30))
    doc.image(_JPEG_8x4)
    doc.bullets(["a", "b"])
    path = doc.save(os.path.join(d, "p.docx"))
    pd = pydocx.Document(path)
    heads = [p.text for p in pd.paragraphs if p.style.name.startswith("Heading")]
    assert heads == ["H1", "H2"], heads
    assert len(pd.tables) == 2 and len(pd.inline_shapes) == 2
    assert pd.core_properties.title == "T" and pd.core_properties.author == ""
    shutil.rmtree(d)


# ----------------------------------------------------------------------------- markdown
def _md_doc(md, **kw):
    d = _tmp()
    doc = R.Docx()
    heads = R.add_markdown(doc, md, **kw)
    path = doc.save(os.path.join(d, "m.docx"))
    xml, _ = _check_package(path)
    return R.inspect_docx(path), xml, heads, d


@test
def markdown_converter_handles_the_tools_own_conventions():
    md = ("# Breach model run: T1\n\nND Dam Breach Runner 2.0.0 · policy · finished x\n\n"
          "## Result\n\n| | |\n|---|---|\n| Scenario | sunny day |\n| Peak | **22,009 cfs** (piping) |\n\n"
          "## Check these\n\n- first `code_span` here\n- second with **bold**\n\n"
          "## Downstream extent (ND 5.1.1(f))\n\n"
          "| Run | Size (mi) | Cells | Over |\n|---|---|---|---|\n| 1 | 2.07 x 1.03 | 8,778 | none |\n"
          "| 2 | 3 x 2 | 12,000 | E, N |\n\n"
          "### Deeper\n\n#### Too deep\n\nLast line.\n")
    info, xml, heads, d = _md_doc(md)
    assert heads == [(1, "Breach model run: T1"), (2, "Result"), (2, "Check these"),
                     (2, "Downstream extent (ND 5.1.1(f))"), (3, "Deeper")], heads
    assert info["tables"] == 2
    t = info["text"]
    assert "sunny day" in t and "22,009 cfs" in t and "code_span" in t and "**" not in t
    assert "| Run |" not in t and "---" not in t          # separators dropped
    assert "Too deep" in t                                 # level 4 becomes a bold paragraph
    assert xml.count("<w:b/>") >= 3                        # bold runs
    assert 'w:styleId' not in xml and 'w:rStyle w:val="Code"' in xml
    # empty header row: label / value table without a header row and a shaded first column
    assert "<w:tblHeader/>" in xml                         # the second table has a header
    assert xml.count("<w:tblHeader/>") == 1
    assert "F2F2F2" in xml
    # bullets under 'Check these' are shaded like the HTML report's warning list
    assert 'w:fill="FFF6E0"' in xml
    # numeric columns right-aligned
    assert '<w:jc w:val="right"/>' in xml
    shutil.rmtree(d)
    # skip_title and shift
    info, xml, heads, d = _md_doc(md, skip_title=True, heading_shift=1)
    assert heads[0] == (3, "Result") and (4, "Deeper") not in heads
    assert "Breach model run" not in info["text"]
    shutil.rmtree(d)


@test
def markdown_converter_lists_code_quotes_and_inline_edge_cases():
    md = ("- a\n  - nested\n    - deeper\n- b\n\n1. one\n2. two\n\nparagraph\n\n1. again\n\n"
          "```\ncode | not a table\n# not a heading\n```\n\n> quoted **text**\n\n---\n\n"
          "path /tmp/a_b_c/hmax_ft.tif and *emph* and _emph2_ and 2 * 3 * 4 and \\*lit\\*\n\n"
          "| A | B |\n|:--|--:|\n| x \\| y | 1 |\n\n[link](http://x.y/z) end\n")
    info, xml, heads, d = _md_doc(md)
    t = info["text"]
    assert "code | not a table" in t and "# not a heading" in t
    assert "a_b_c/hmax_ft.tif" in t                        # underscores inside words stay
    assert "emph2" in t and "_emph2_" not in t and "*emph*" not in t
    assert "2 * 3 * 4" in t and "*lit*" in t
    assert "x | y" in t and "link (http://x.y/z) end" in t
    assert xml.count('<w:ilvl w:val="2"/>') == 1           # three list levels
    assert len(set(re.findall(r'<w:numId w:val="(\d+)"/>', xml)) - {"1"}) == 2
    assert info["tables"] == 1 and heads == []
    assert "<w:i/>" in xml
    shutil.rmtree(d)
    # standard markdown: lines of one paragraph joined on request
    info, _x, _h, d = _md_doc("line one\nline two\n\nnext", join_lines=True)
    assert "line one line two" in info["text"]
    info2, _x, _h, d2 = _md_doc("line one\nline two\n\nnext")
    assert "line one\nline two" in info2["text"]
    # nothing at all, and text that is only markup characters
    for src in ("", "\n\n", "|", "#", "**", "- ", "|---|"):
        _md_doc(src)
    for x in (d, d2):
        shutil.rmtree(x)


@test
def inline_markdown_runs():
    runs = R.parse_inline("a **b** `c` *d* e_f_g")
    assert [(r["text"], sorted(k for k in r if k != "text")) for r in runs] == [
        ("a ", []), ("b", ["b"]), (" ", []), ("c", ["code"]), (" ", []), ("d", ["i"]),
        (" e_f_g", [])], runs
    assert R.parse_inline("") == []
    assert R.label_of("total_cfs") == "Total (cfs)" and R.label_of("max_depth_ft") == "Maximum depth (ft)"
    assert R.fmt_val(12345.678) == "12,346" and R.fmt_val(0.5000) == "0.5"
    assert R.fmt_val(None) == u"—" and R.fmt_val(True) == "Yes"


# ----------------------------------------------------------------------------- SVG -> PNG
_SVG = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200" width="400" height="200">'
        '<rect width="100%" height="100%" fill="#fff"/><circle cx="100" cy="100" r="50" '
        'fill="red"/><text x="200" y="100" font-size="20">Hello</text></svg>')


@test
def svg_to_png_uses_whatever_backend_exists_and_returns_none_otherwise():
    d = _tmp()
    svg = os.path.join(d, "f.svg")
    open(svg, "w").write(_SVG)
    assert R.svg_size(svg) == (400.0, 200.0)
    open(os.path.join(d, "g.svg"), "w").write('<svg width="30" height="10"></svg>')
    assert R.svg_size(os.path.join(d, "g.svg")) == (30.0, 10.0)
    assert R.svg_size(os.path.join(d, "missing.svg")) is None
    png = os.path.join(d, "f.png")
    errs = []
    assert R.svg_to_png(svg, png, 300, backends=["nothing-like-this"], errors=errs) is None
    assert not os.path.exists(png)
    # a converter that is not on this machine fails cleanly (and says why)
    errs = []
    saved = os.environ.get("PATH", "")
    os.environ["PATH"] = ""
    try:
        assert R.svg_to_png(svg, png, 300, backends=["rsvg-convert", "inkscape"],
                            errors=errs) is None
    finally:
        os.environ["PATH"] = saved
    assert len(errs) == 2 and not os.path.exists(png)
    avail = R.available_svg_backends()
    got = R.svg_to_png(svg, png, 640)
    if not avail:
        assert got is None
    else:
        assert got == png, R.svg_to_png(svg, png, 640, errors=errs) or errs
        fmt, w, h, _dx, _dy = R.image_info(open(png, "rb").read())
        assert fmt == "png" and w == 640 and h == 320, (w, h)
        assert R.LAST_BACKEND in avail
    # a damaged SVG never raises
    bad = os.path.join(d, "bad.svg")
    open(bad, "w").write("<svg><<<")
    R.svg_to_png(bad, os.path.join(d, "bad.png"), 200, backends=["cairosvg", "rsvg-convert"])
    shutil.rmtree(d)


# ----------------------------------------------------------------------------- a run folder
_CACHE = {}


def _run_folder(flood=True):
    """A dam folder written by the real workflow (the LISFLOOD-FP call is stubbed)."""
    key = "flood" if flood else "sunny"
    if key in _CACHE:
        return _CACHE[key]
    from nd_breach_runner.core import breach as br, flood as floodmod, impacts, workflow
    from nd_breach_runner.core import lisflood as lf
    from nd_breach_runner.tests.synthetic import SyntheticTerrain, npy_writer

    def fake_run(wd, par, st, progress=None, cancel=None, log=None, name="run"):
        return 90.0

    def fake_read(wd, dem, st, hydrograph, name="run"):
        X, Y = np.meshgrid(dem.x_centers(), dem.y_centers())
        depth = np.where((X > 0) & (np.abs(Y) < 400),
                         2.0 * np.exp(-X / 1500.0) * np.exp(-(np.abs(Y) / 300.0) ** 2), np.nan)
        depth = np.where(depth > 0.01, depth, np.nan)
        vel = np.where(np.isfinite(depth), 1.5, np.nan)
        arr = np.where(np.isfinite(depth), X / 1.5 / 3600.0 + 6.3, np.nan)
        wse = np.where(np.isfinite(depth), dem.data + np.nan_to_num(depth), np.nan)
        stats = {"cells": int(dem.data.size), "cell_size_m": dem.dx, "sim_time_hr": 4.0,
                 "max_depth_ft": 6.5, "max_velocity_fps": 4.9, "flooded_acres": 300.0,
                 "peak_inflow_cfs": 16000.0, "max_volume_error_pct": 0.12,
                 "worst_interval_error_pct": 0.4, "injected_acft": 499.0, "in_domain_acft": 350.0}
        return lf.Envelopes(dem.like(depth, np.nan), dem.like(vel, np.nan),
                            dem.like(arr, np.nan), dem.like(wse, np.nan), stats, [], {})
    d = br.DamData(dam_id="ND00123", dam_height_ft=25.0, storage_top_acft=500.0,
                   top_of_dam_ft=1850.0, surface_area_top_ac=60.0, crest_width_ft=12,
                   us_slope_h=3, ds_slope_h=2.5, name="Test Creek Dam", crest_length_ft=800.0)
    st = workflow.RunSettings(cell_size_m=50.0, planning_cell_m=90.0, initial_length_mi=1.0, half_width_mi=0.5,
                              expand_step_mi=1.0, max_iterations=3, lisflood_exe=sys.executable)
    rec = impacts.Receptors()
    rec.add_points([(500, 0), (900, 50), (1400, -30), (2500, 20)], ids=["h1", "h2", "h3", "h4"])
    st.receptors = rec
    if flood:
        st.scenario = "flood"
        st.flood = floodmod.FloodInputs.triangle(
            6000.0, 6.0, 30.0, spillway=floodmod.Spillway.from_weir(1846.0, 20.0),
            normal_pool_ft=1846.0)
    root = os.path.join(_tmp(), "dams")
    folder = os.path.join(root, "ND00123")
    old = (lf.run, lf.read_results)
    lf.run, lf.read_results = fake_run, fake_read
    try:
        res = workflow.run_dam(d, (0.0, 0.0), st, SyntheticTerrain(), npy_writer, folder)
    finally:
        lf.run, lf.read_results = old
    assert res.summary["breach_occurs"]
    _CACHE[key] = folder
    return folder


def _add_extras(folder):
    """What other modules and tools write: sensitivity, consequences, manifest, cascade, gauge."""
    from nd_breach_runner.core import sensitivity as sens
    so = sens.SensitivityResult()
    so.folder = os.path.join(folder, "Sensitivity")
    os.makedirs(so.folder)
    so.rows = [{"case": "Base", "group": "base", "method": "piping", "peak_cfs": 16175,
                "flooded_ac": 312.4, "max_depth_ft": 6.5, "structures_flooded": 3},
               {"case": "Width x0.7", "group": "breach width", "method": "piping",
                "peak_cfs": 12100, "flooded_ac": 280.0, "max_depth_ft": 5.9,
                "structures_flooded": 0}]
    so.flags = ["Flooded area: 280 ac in the lowest case, 340 ac in the highest.",
                "CHANGES THE ANSWER: some cases flood no structure and others do."]
    so.warnings = ["Time x2: the dam does not fail in this case."]
    so.envelope = {"Sensitivity Depth (Low).tif": "x"}
    sens.write_outputs(so, None)
    open(os.path.join(folder, "consequences_summary.csv"), "w").write(
        "Reach,Structures,Population at risk,Loss ($)\nA,3,12,450000\nB,1,4,90000\n")
    open(os.path.join(folder, "Consequences_Notes.md"), "w").write(
        "# Consequences\n\n## Basis\n\n- Population from **2020 census** blocks\n\n| Reach | PAR |\n"
        "|---|---:|\n| A | 12 |\n")
    json.dump({"tool_versions": {"nd_breach_runner": "2.0.0", "python": "3.11"},
               "inputs": [{"path": "dem.tif", "sha256": "a" * 64},
                          {"path": "s.gpkg", "sha256": "b" * 64}],
               "settings": {"cell_size_m": 25, "nested": {"a": 1}}},
              open(os.path.join(folder, "_breach_model", "run_manifest.json"), "w"))
    parent = os.path.dirname(folder)
    os.makedirs(os.path.join(parent, "Cascade_ND00123"))
    open(os.path.join(parent, "Cascade_ND00123", "Cascade_Report.md"), "w").write(
        "# Cascade failure\n\nDams upstream to downstream: ND00123, ND00124\n\n## Decisions\n\n"
        "- ND00124: FAILS (assumed to fail)\n")
    open(os.path.join(folder, "_breach_model", "gauge_bridge.csv"), "w").write(
        "time_hr,discharge_cfs\n0,0\n1,5000\n2,3000\n")


def _copy_folder(flood=True, extras=False):
    src = _run_folder(flood)
    root = _tmp()
    dst = os.path.join(root, "dams", os.path.basename(src))
    shutil.copytree(src, dst)
    if extras:
        _add_extras(dst)
    return dst


def _status(info_text, item):
    """The checklist status of the row whose item text starts with `item`."""
    m = re.search(re.escape(item) + r"[^\n]*\n(Yes|Check|No|N/A)\n", info_text)
    return m.group(1) if m else None


@test
def submittal_from_a_real_run_folder_has_all_sections():
    folder = _copy_folder(True, extras=True)
    out = os.path.join(folder, "Submittal_ND00123.docx")
    opt = R.SubmittalOptions(engineer="Jane Engineer, PE", firm="Prairie Engineering",
                             project="Test Creek EAP", include_run_report=True,
                             created=__import__("datetime").datetime(2026, 9, 29, 12, 0, 0))
    res = R.build_submittal(folder, "", out, opt)
    # without any SVG converter the figures are listed by name and the report says so (the only
    # warning then; see the figures check below)
    warnings = [w for w in res.warnings
                if R.available_svg_backends() or not w.startswith("No SVG converter")]
    assert res.path == out and os.path.exists(out) and not warnings, res.warnings
    xml, names = _check_package(out)
    info = R.inspect_docx(out)
    titles = [t for _n, t in res.sections]
    assert titles == ["Summary of results", "Dam and reservoir data",
                      "Breach parameters and method", "Flood scenario", "Hydraulic model",
                      "Results", "Receptors and impacts",
                      "Sensitivity, cascade and probabilistic results",
                      "Assumptions and limitations", "Reproducibility", "Appendix"], titles
    assert [n for n, _t in res.sections] == [str(i) for i in range(1, 12)]
    t = info["text"]
    rep = json.load(open(os.path.join(folder, "_breach_model", "run_report.json")))
    peak = "{:,.0f}".format(rep["summary"]["peak_breach_cfs"])
    assert peak in t and "Test Creek Dam" in t and "Jane Engineer, PE" in t
    assert "Prairie Engineering" in t and "Test Creek EAP" in t
    # sections: the pieces named in the brief
    for needle in ("Piping failure", "Notes and warnings recorded by the run",
                   "Baseline and breach outflow", "Incremental effect of the breach",
                   "Mass balance", "Flooded area by depth class", "Arrival times",
                   "Conditions at flooded receptors", "Receptors reached first",
                   "Sensitivity of the breach model", "CHANGES THE ANSWER",
                   "Cascade failure", "Assumptions made by the tool",
                   "To be verified before submittal", "Run manifest", "Run settings",
                   "Reviewer checklist", "Files in the dam folder", "Run report as written",
                   "consequences_summary.csv", "Population from 2020 census blocks",
                   "a" * 64, "Reviewer OK / comment", "Peak discharge (cfs)"):
        assert needle in t, needle
    # the checklist is pre-filled from the folder
    for item, want in (("EAP inundation map produced", "Yes"),
                       ("Depth, velocity, arrival-time and WSE rasters produced", "Yes"),
                       ("Incremental consequences computed", "Yes"),
                       ("Consequences documented", "Yes"),
                       ("Run manifest with software versions", "Yes"),
                       ("Sensitivity to the breach parameters evaluated", "Yes"),
                       ("Model mass balance within 1%", "Yes"),
                       ("Run warnings resolved or explained", "Check"),
                       ("Hazard classification made by the engineer", "Check"),
                       ("Cross-check with a second model", "No")):
        assert _status(t, item) == want, (item, _status(t, item))
    # the run stopped growing the domain where no receptor lay ahead: a check, not a pass
    assert _status(t, "Downstream extent criterion (5.1.1(f)) met") == "Check"
    # the run's own markdown was inserted, tables and all
    assert info["tables"] >= 25 and info["images"] >= 0
    # figures: embedded when a converter exists, otherwise listed by file name
    figs = {f["name"]: f for f in res.figures}
    assert set(figs) == {"EAP inundation map", "Outflow hydrograph"}
    if R.available_svg_backends():
        assert all(f["embedded"] for f in figs.values()) and info["images"] == 2, res.figures
        assert res.svg_backend in R.available_svg_backends()
        assert "Figure 1" in t and "Figure 2" in t
    else:
        assert not any(f["embedded"] for f in figs.values())
        assert "not embedded" in t and "EAP Inundation Map (preview).svg" in t
    assert "run_manifest.json" not in res.missing and "EAP Inundation Map.svg" not in res.missing
    # the report never wrote into the dam folder other than the output file itself
    assert not [n for n in os.listdir(folder) if n.startswith("ndbr")]
    shutil.rmtree(os.path.dirname(os.path.dirname(folder)))


@test
def submittal_options_switch_sections_off_and_leave_no_gaps():
    folder = _copy_folder(True, extras=True)
    out = os.path.join(folder, "s.docx")
    res = R.build_submittal(folder, "ND00123", out, {
        "include_figures": False, "include_sensitivity": False, "include_consequences": False,
        "include_reproducibility": False, "include_checklist": False, "include_file_list": True,
        "include_assumptions": False, "include_toc": False})
    info = R.inspect_docx(out)
    assert info["images"] == 0 and not res.figures
    names = [t for _n, t in res.sections]
    assert "Sensitivity, cascade and probabilistic results" not in names
    assert "Reproducibility" not in names and "Assumptions and limitations" not in names
    assert [n for n, _t in res.sections] == [str(i) for i in range(1, len(names) + 1)]
    t = info["text"]
    assert "Reviewer checklist" not in t and "consequences_summary" not in t.split(
        "Files in the dam folder")[0]
    assert "TOC" not in zipfile.ZipFile(out).read("word/document.xml").decode()
    assert "Files in the dam folder" in t
    # the checklist marks a switched-off consequences as absent
    try:
        R.SubmittalOptions(nonsense=1)
        assert False
    except TypeError:
        pass
    assert R.SubmittalOptions.from_any(None).title.startswith("Dam Breach")
    shutil.rmtree(os.path.dirname(os.path.dirname(folder)))


@test
def submittal_sunny_day_and_no_breach_folders():
    folder = _copy_folder(False)
    out = os.path.join(folder, "s.docx")
    res = R.build_submittal(folder, "", out, {"include_figures": False})
    t = R.inspect_docx(out)["text"]
    assert "sunny-day (no inflow) breach" in t
    assert _status(t, "Without-breach (baseline) flood modelled") == "N/A"
    assert _status(t, "Incremental consequences computed") == "N/A"
    assert "Baseline and breach outflow" not in t
    assert not res.warnings, res.warnings
    shutil.rmtree(os.path.dirname(os.path.dirname(folder)))
    # what the workflow writes when the dam does not fail in the flood
    d = _tmp()
    mdl = os.path.join(d, "_breach_model")
    os.makedirs(mdl)
    rep = {"tool": "ND Dam Breach Runner 2.0.0", "scenario": "flood",
           "dam": {"dam_id": "X1", "name": "", "dam_height_ft": 20.0},
           "settings": {"tier": "lisflood"}, "flood": {"failed": False},
           "summary": {"dam_id": "X1", "final_model": "none", "breach_occurs": False,
                       "flooded_acres": 0},
           "warnings": ["The pool never reaches the trigger."]}
    json.dump(rep, open(os.path.join(mdl, "run_report.json"), "w"))
    open(os.path.join(mdl, "run_report.md"), "w").write(
        "# Breach model run: X1\n\nThe dam does not fail in this flood.\n\n- The pool never "
        "reaches the trigger.\n")
    out = os.path.join(d, "n.docx")
    res = R.build_submittal(d, "X1", out)
    info = R.inspect_docx(out)
    _check_package(out)
    assert "The dam does not fail in this flood" in info["text"]
    assert _status(info["text"], "Downstream extent criterion (5.1.1(f)) met") == "No"
    shutil.rmtree(d)


@test
def submittal_never_fails_on_missing_or_damaged_files():
    d = _tmp()
    cases = []
    # 1 an empty folder
    e = os.path.join(d, "empty")
    os.makedirs(e)
    cases.append((e, ""))
    # 2 a folder that does not exist
    cases.append((os.path.join(d, "nope"), "GONE"))
    # 3 garbage where the JSON should be, and a report.md only
    g = os.path.join(d, "garbage")
    os.makedirs(os.path.join(g, "_breach_model"))
    open(os.path.join(g, "_breach_model", "run_report.json"), "w").write("{not json")
    open(os.path.join(g, "_breach_model", "run_report.md"), "w").write("# Run\n\n- ok\n")
    cases.append((g, "G1"))
    # 4 valid JSON of the wrong shapes
    for k, payload in enumerate(("[]", "5", "{}", '{"summary": 3, "hydrograph": [1], '
                                 '"settings": "x", "dam": null, "warnings": "oops", '
                                 '"lisflood_runs": [{}, null, 5]}',
                                 '{"summary": {"dam_id": "Z", "domain_mi": [1], '
                                 '"incremental_acres": 4}, "flood": 1, "receptors": {"structure": 2}}')):
        f = os.path.join(d, "shape%d" % k)
        os.makedirs(os.path.join(f, "_breach_model"))
        open(os.path.join(f, "_breach_model", "run_report.json"), "w").write(payload)
        cases.append((f, "S%d" % k))
    # 5 damaged companion files
    b = os.path.join(d, "broken")
    os.makedirs(os.path.join(b, "_breach_model"))
    os.makedirs(os.path.join(b, "Sensitivity"))
    open(os.path.join(b, "Sensitivity", "Sensitivity_Summary.csv"), "wb").write(b"\xff\xfe\x00bad")
    open(os.path.join(b, "Sensitivity", "Sensitivity_Report.html"), "w").write("<h2>What")
    for name in ("Flood Polygons (Max Depth).geojson", "Arrival Zones.geojson",
                 "run_manifest.json", "consequences_x.csv", "consequences_y.md"):
        open(os.path.join(b, name), "wb").write(b"\x00\x01garbage{[")
    open(os.path.join(b, "_breach_model", "receptors_flooded.csv"), "w").write("a,b\n1\n")
    open(os.path.join(b, "_breach_model", "breach_hydrograph.csv"), "w").write("x\n")
    open(os.path.join(b, "EAP Inundation Map.svg"), "w").write("<svg><<<")
    open(os.path.join(b, "hydrograph.png"), "wb").write(b"not a png")
    open(os.path.join(b, "_breach_model", "hydrograph.png"), "wb").write(b"not a png")
    open(os.path.join(b, "_breach_model", "hydrograph.svg"), "w").write("")
    open(os.path.join(b, "_breach_model", "gauge_a.csv"), "w").write("time_hr,discharge_cfs\n")
    cases.append((b, "B1"))
    for k, (folder, dam_id) in enumerate(cases):
        out = os.path.join(d, "out%d.docx" % k)
        res = R.build_submittal(folder, dam_id, out, {"svg_backends": ["cairosvg"]})
        xml, _n = _check_package(out)
        info = R.inspect_docx(out)
        titles = [t for _n2, t in res.sections]
        assert "Summary of results" in titles and "Assumptions and limitations" in titles \
            and "Appendix" in titles, (folder, titles)
        assert "Reviewer checklist" in info["text"]
    # the nonexistent folder is reported, not raised
    out = os.path.join(d, "out1.docx")
    assert R.build_submittal(cases[1][0], "GONE", out).warnings
    # an unwritable output path is the one thing that raises
    try:
        R.build_submittal(e, "", os.path.join(d, "out0.docx", "sub", "x.docx"))
        assert False
    except OSError:
        pass
    shutil.rmtree(d)


@test
def submittal_reads_the_shared_files_by_name_case_insensitively():
    folder = _copy_folder(True, extras=True)
    mdl = os.path.join(folder, "_breach_model")
    os.rename(os.path.join(mdl, "run_manifest.json"), os.path.join(folder, "RUN_MANIFEST.JSON"))
    os.makedirs(os.path.join(folder, "Consequences"))
    open(os.path.join(folder, "Consequences", "table.csv"), "w").write("A,B\n1,2\n")
    open(os.path.join(mdl, "probabilistic_summary.csv"), "w").write("p,depth\n0.5,3.1\n")
    out = os.path.join(folder, "c.docx")
    res = R.build_submittal(folder, "", out, {"include_figures": False, "max_table_rows": 3})
    t = R.inspect_docx(out)["text"]
    assert "Run manifest" in t and "a" * 64 in t
    assert "Probabilistic and uncertainty results" in t and "0.5" in t
    assert "table.csv" in t
    assert "Only the first 3 rows" in t or "Only the first" not in t     # truncation note works
    assert not res.warnings
    shutil.rmtree(os.path.dirname(os.path.dirname(folder)))


# ----------------------------------------------------------------------------- the algorithm
@test
def submittal_algorithm_runs_through_the_fake_harness():
    from nd_breach_runner.tests import fake_qgis as fq
    fq.install()
    import importlib
    import nd_breach_runner.algorithms as algs
    importlib.reload(algs)
    import nd_breach_runner.algorithms_tools as tools
    importlib.reload(tools)
    import nd_breach_runner.algorithms_report as rep
    importlib.reload(rep)
    assert rep.REPORT_ALGORITHMS == (rep.SubmittalReportAlgorithm,)
    assert issubclass(rep.SubmittalReportAlgorithm, tools._ToolBase)
    a = rep.SubmittalReportAlgorithm()
    a.initAlgorithm()
    assert a.name() == "submittal" and a.displayName() == "Submittal report (Word)"
    assert a.shortHelpString() and a.group() == "Data preparation and tools"
    for p in ("DAM_FOLDER", "DAM_ID", "DAM_NAME", "TITLE", "ENGINEER", "FIRM", "LOGO",
              "INC_FIGURES", "INC_SENSITIVITY", "INC_CONSEQUENCES", "INC_REPRO", "INC_CHECKLIST",
              "INC_FILES", "OUTPUT"):
        assert p in a.params, p
    assert {"REPORT", "RESULT"} <= set(a.outputs)
    folder = _copy_folder(True, extras=True)
    d = os.path.dirname(os.path.dirname(folder))
    logo = os.path.join(d, "logo.png")
    open(logo, "wb").write(_png(60, 30))
    out = os.path.join(d, "reports", "sub.docx")
    fb = fq.QgsProcessingFeedback()
    ctx = fq.QgsProcessingContext()
    r = a.processAlgorithm({"DAM_FOLDER": folder, "DAM_ID": "", "DAM_NAME": "Test Creek Dam",
                            "TITLE": "Dam Breach Report", "ENGINEER": "J. Engineer, PE",
                            "FIRM": "Prairie Engineering", "LOGO": logo, "INC_FIGURES": False,
                            "INC_SENSITIVITY": True, "INC_CONSEQUENCES": True,
                            "INC_REPRO": True, "INC_CHECKLIST": True, "INC_FILES": False,
                            "OUTPUT": out}, ctx, fb)
    assert r["REPORT"] == out and os.path.exists(out)
    _check_package(out)
    info = R.inspect_docx(out)
    assert info["images"] == 1                                  # the logo
    assert "J. Engineer, PE" in info["text"] and "Dam Breach Report" in info["text"]
    assert "Files in the dam folder" not in info["text"]
    assert "Sensitivity of the breach model" in info["text"]
    assert any("11" in m and "section" in m.lower() for m in fb.info), fb.info
    assert "sections" in r["RESULT"]
    # the dam ID falls back to the folder name; a missing folder is an error
    assert "ND00123" in info["text"]
    try:
        a.processAlgorithm({"DAM_FOLDER": os.path.join(d, "nope"), "OUTPUT": out}, ctx, fb)
        assert False
    except fq.QgsProcessingException as e:
        assert "folder" in str(e).lower()
    try:
        a.processAlgorithm({"DAM_FOLDER": folder, "OUTPUT": ""}, ctx, fb)
        assert False
    except fq.QgsProcessingException:
        pass
    shutil.rmtree(d)
